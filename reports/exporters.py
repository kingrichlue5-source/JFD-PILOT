"""
Export utilities for HMIS reports: CSV and formatted XLSX.
"""
import csv
import io
from decimal import Decimal

from django.http import HttpResponse


def export_csv(indicators, filename='hmis_report'):
    """Export HMIS indicators as a CSV file."""
    response = HttpResponse(content_type='text/csv')
    response['Content-Disposition'] = f'attachment; filename="{filename}.csv"'

    writer = csv.writer(response)
    writer.writerow(['Indicator', 'Value'])

    for key, value in indicators.items():
        label = key.replace('_', ' ').title()
        writer.writerow([label, value])

    return response


def export_xlsx(hmis_data, filename='hmis_report'):
    """Export HMIS report as a formatted XLSX file with hospital header, colored sections, and chart."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.chart import BarChart, Reference
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = 'HMIS Monthly Report'

    # ── Color scheme ──────────────────────────────────────────────────
    red_fill = PatternFill(start_color='DC2626', end_color='DC2626', fill_type='solid')
    blue_fill = PatternFill(start_color='1E40AF', end_color='1E40AF', fill_type='solid')
    light_blue_fill = PatternFill(start_color='DBEAFE', end_color='DBEAFE', fill_type='solid')
    light_gray_fill = PatternFill(start_color='F3F4F6', end_color='F3F4F6', fill_type='solid')
    white_font = Font(color='FFFFFF', bold=True, size=12)
    header_font = Font(color='FFFFFF', bold=True, size=11)
    title_font = Font(bold=True, size=16, color='1E3A8A')
    subtitle_font = Font(bold=True, size=12, color='4B5563')
    border = Border(
        left=Side(style='thin', color='D1D5DB'),
        right=Side(style='thin', color='D1D5DB'),
        top=Side(style='thin', color='D1D5DB'),
        bottom=Side(style='thin', color='D1D5DB'),
    )

    # ── Column widths ─────────────────────────────────────────────────
    ws.column_dimensions['A'].width = 40
    ws.column_dimensions['B'].width = 20
    ws.column_dimensions['C'].width = 5
    ws.column_dimensions['D'].width = 40
    ws.column_dimensions['E'].width = 20

    # ── Hospital Header ───────────────────────────────────────────────
    ws.merge_cells('A1:E1')
    cell = ws['A1']
    cell.value = 'JACKSON F. DOE MEMORIAL REGIONAL REFERRAL HOSPITAL'
    cell.font = Font(bold=True, size=18, color='1E3A8A')
    cell.alignment = Alignment(horizontal='center')

    ws.merge_cells('A2:E2')
    cell = ws['A2']
    cell.value = 'Hospital Management Information System (HMIS) Monthly Report'
    cell.font = subtitle_font
    cell.alignment = Alignment(horizontal='center')

    ws.merge_cells('A3:E3')
    cell = ws['A3']
    cell.value = f"Reporting Period: {hmis_data.get('period', 'N/A')}"
    cell.font = Font(bold=True, size=11, color='6B7280')
    cell.alignment = Alignment(horizontal='center')

    ws.merge_cells('A4:E4')
    cell = ws['A4']
    cell.value = f"Organization: {hmis_data.get('org_unit', 'Jackson F. Doe Memorial Regional Referral Hospital')}"
    cell.font = Font(size=10, color='6B7280')
    cell.alignment = Alignment(horizontal='center')

    row = 6

    # ── Attendance Section ────────────────────────────────────────────
    ws.merge_cells(f'A{row}:B{row}')
    cell = ws[f'A{row}']
    cell.value = 'ATTENDANCE & PATIENT FLOW'
    cell.font = white_font
    cell.fill = blue_fill
    ws[f'B{row}'].fill = blue_fill
    row += 1

    indicators = hmis_data.get('indicators', {})
    attendance_items = [
        ('Total Attendance (Unique Patients)', indicators.get('total_attendance', 0)),
        ('OPD Attendance', indicators.get('opd_attendance', 0)),
        ('ER Attendance', indicators.get('er_attendance', 0)),
        ('IPD Admissions', indicators.get('ipd_admissions', 0)),
    ]

    for label, value in attendance_items:
        ws[f'A{row}'] = label
        ws[f'A{row}'].font = Font(bold=True)
        ws[f'A{row}'].border = border
        ws[f'B{row}'] = value
        ws[f'B{row}'].alignment = Alignment(horizontal='center')
        ws[f'B{row}'].border = border
        if row % 2 == 0:
            ws[f'A{row}'].fill = light_gray_fill
            ws[f'B{row}'].fill = light_gray_fill
        row += 1

    row += 1

    # ── Maternal & Newborn Section ────────────────────────────────────
    ws.merge_cells(f'A{row}:B{row}')
    cell = ws[f'A{row}']
    cell.value = 'MATERNAL & NEWBORN HEALTH'
    cell.font = white_font
    cell.fill = red_fill
    ws[f'B{row}'].fill = red_fill
    row += 1

    maternal_items = [
        ('Maternal Deliveries', indicators.get('maternal_deliveries', 0)),
        ('Newborn Deliveries', indicators.get('newborn_deliveries', 0)),
        ('Mortality Count', indicators.get('mortality_count', 0)),
    ]

    for label, value in maternal_items:
        ws[f'A{row}'] = label
        ws[f'A{row}'].font = Font(bold=True)
        ws[f'A{row}'].border = border
        ws[f'B{row}'] = value
        ws[f'B{row}'].alignment = Alignment(horizontal='center')
        ws[f'B{row}'].border = border
        if row % 2 == 0:
            ws[f'A{row}'].fill = light_gray_fill
            ws[f'B{row}'].fill = light_gray_fill
        row += 1

    row += 1

    # ── Services Section ──────────────────────────────────────────────
    ws.merge_cells(f'A{row}:B{row}')
    cell = ws[f'A{row}']
    cell.value = 'CLINICAL SERVICES'
    cell.font = white_font
    cell.fill = blue_fill
    ws[f'B{row}'].fill = blue_fill
    row += 1

    services_items = [
        ('Laboratory Orders', indicators.get('lab_orders', 0)),
        ('Radiology Orders', indicators.get('radiology_orders', 0)),
        ('Prescriptions Issued', indicators.get('prescriptions_issued', 0)),
        ('Medications Dispensed', indicators.get('medications_dispensed', 0)),
    ]

    for label, value in services_items:
        ws[f'A{row}'] = label
        ws[f'A{row}'].font = Font(bold=True)
        ws[f'A{row}'].border = border
        ws[f'B{row}'] = value
        ws[f'B{row}'].alignment = Alignment(horizontal='center')
        ws[f'B{row}'].border = border
        if row % 2 == 0:
            ws[f'A{row}'].fill = light_gray_fill
            ws[f'B{row}'].fill = light_gray_fill
        row += 1

    row += 1

    # ── Revenue Section ───────────────────────────────────────────────
    ws.merge_cells(f'A{row}:B{row}')
    cell = ws[f'A{row}']
    cell.value = 'FINANCIAL SUMMARY'
    cell.font = white_font
    cell.fill = red_fill
    ws[f'B{row}'].fill = red_fill
    row += 1

    ws[f'A{row}'] = 'Revenue Collected'
    ws[f'A{row}'].font = Font(bold=True)
    ws[f'A{row}'].border = border
    ws[f'B{row}'] = indicators.get('revenue_collected', '0.00')
    ws[f'B{row}'].alignment = Alignment(horizontal='center')
    ws[f'B{row}'].border = border
    row += 2

    # ── Top Diagnoses Section ─────────────────────────────────────────
    top_diag = hmis_data.get('top_diagnoses', [])
    if top_diag:
        ws.merge_cells(f'A{row}:B{row}')
        cell = ws[f'A{row}']
        cell.value = 'TOP 10 DIAGNOSES'
        cell.font = white_font
        cell.fill = blue_fill
        ws[f'B{row}'].fill = blue_fill
        row += 1

        # Table headers
        ws[f'A{row}'] = 'Diagnosis'
        ws[f'A{row}'].font = Font(bold=True, color='1E3A8A')
        ws[f'A{row}'].border = border
        ws[f'B{row}'] = 'Count'
        ws[f'B{row}'].font = Font(bold=True, color='1E3A8A')
        ws[f'B{row}'].alignment = Alignment(horizontal='center')
        ws[f'B{row}'].border = border
        row += 1

        chart_start_row = row
        for i, diag in enumerate(top_diag[:10], 1):
            desc = diag.get('diagnosis_description', 'N/A')
            code = diag.get('diagnosis_code', '')
            label = f"{code}: {desc}" if code else desc
            if len(label) > 60:
                label = label[:57] + '...'
            ws[f'A{row}'] = label
            ws[f'A{row}'].font = Font(size=10)
            ws[f'A{row}'].border = border
            ws[f'B{row}'] = diag.get('count', 0)
            ws[f'B{row}'].alignment = Alignment(horizontal='center')
            ws[f'B{row}'].border = border
            if row % 2 == 0:
                ws[f'A{row}'].fill = light_gray_fill
                ws[f'B{row}'].fill = light_gray_fill
            row += 1

        chart_end_row = row - 1

        # ── Bar Chart for Top Diagnoses ───────────────────────────────
        if chart_end_row >= chart_start_row:
            chart = BarChart()
            chart.type = 'bar'
            chart.title = 'Top 10 Diagnoses'
            chart.y_axis.title = 'Count'
            chart.x_axis.title = 'Diagnosis'
            chart.style = 10
            chart.width = 30
            chart.height = 15

            data_ref = Reference(ws, min_col=2, min_row=chart_start_row - 1, max_row=chart_end_row)
            cats_ref = Reference(ws, min_col=1, min_row=chart_start_row, max_row=chart_end_row)
            chart.add_data(data_ref, titles_from_data=True)
            chart.set_categories(cats_ref)
            chart.shape = 4

            ws.add_chart(chart, f'D6')

    # ── Footer ────────────────────────────────────────────────────────
    row += 2
    ws.merge_cells(f'A{row}:B{row}')
    ws[f'A{row}'] = 'Generated by JFD Hospital Management Information System'
    ws[f'A{row}'].font = Font(italic=True, size=9, color='9CA3AF')

    row += 1
    ws.merge_cells(f'A{row}:B{row}')
    ws[f'A{row}'] = f'Data formatted for DHIS2 compliance | Period: {hmis_data.get("period", "N/A")}'
    ws[f'A{row}'].font = Font(italic=True, size=9, color='9CA3AF')

    # ── Print settings ────────────────────────────────────────────────
    ws.print_area = f'A1:B{row}'
    ws.page_setup.orientation = 'portrait'
    ws.page_setup.fitToWidth = 1

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}.xlsx"'

    wb.save(response)
    return response
